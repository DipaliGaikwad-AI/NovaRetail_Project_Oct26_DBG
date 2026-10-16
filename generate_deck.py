from pathlib import Path
from statistics import mean
import openpyxl
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Pt

ROOT = Path(__file__).resolve().parent
DATA_FILE = ROOT / 'novaretail_dataset.xlsx'
TEMPLATE_FILE = ROOT / 'novaretail_template.pptx'
OUTPUT_FILE = ROOT / 'NovaRetail_QBR_Final.pptx'
ICON_DIR = ROOT / 'icons'

QUARTER_ORDER = ['2025-Q3', '2025-Q4', '2026-Q1', '2026-Q2']
REGION_ORDER = ['North America', 'Europe', 'APAC', 'Latin America']
NAVY = RGBColor(30, 39, 97)
BLUE = RGBColor(66, 103, 178)
TEAL = RGBColor(24, 151, 143)
GOLD = RGBColor(241, 179, 62)
RED = RGBColor(200, 65, 65)
GRAY = RGBColor(91, 107, 140)
WHITE = RGBColor(255, 255, 255)
PALETTE = [BLUE, TEAL, GOLD, RED]


def load_sheet(sheet_name):
    wb = openpyxl.load_workbook(DATA_FILE, data_only=True, read_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    headers = list(rows[0])
    return [dict(zip(headers, row)) for row in rows[1:]]


def index_rows(rows):
    return {(r['Quarter'], r['Region']): r for r in rows}


def pct_change(new, old):
    return ((new / old) - 1) * 100


def shape_by_name(slide, name):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    raise KeyError(f"Shape '{name}' not found")


def remove_shape(shape):
    element = shape._element
    element.getparent().remove(element)


def set_text(slide, name, text, size=None, color=None, bold=None, italic=None,
             bullet_lines=False, margin=0.06, valign=MSO_ANCHOR.TOP):
    shape = shape_by_name(slide, name)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = int(margin * 914400)
    tf.vertical_anchor = valign
    lines = text if isinstance(text, list) else [text]
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        if bullet_lines:
            p.text = f"• {line}"
        p.space_after = Pt(8 if bullet_lines else 2)
        p.alignment = PP_ALIGN.LEFT
        if p.runs:
            run = p.runs[0]
            if size: run.font.size = Pt(size)
            if color: run.font.color.rgb = color
            if bold is not None: run.font.bold = bold
            if italic is not None: run.font.italic = italic
            run.font.name = 'Aptos'
    return shape


def replace_icon(slide, placeholder_name, label_name, icon_file):
    box = shape_by_name(slide, placeholder_name)
    left, top, width, height = box.left, box.top, box.width, box.height
    remove_shape(box)
    remove_shape(shape_by_name(slide, label_name))
    slide.shapes.add_picture(str(ICON_DIR / icon_file), left, top, width, height)


def add_line_chart(slide, placeholder_name, label_name, categories, series, y_title):
    box = shape_by_name(slide, placeholder_name)
    left, top, width, height = box.left, box.top, box.width, box.height
    remove_shape(box)
    remove_shape(shape_by_name(slide, label_name))
    data = ChartData()
    data.categories = categories
    for name, values in series.items():
        data.add_series(name, values)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, left, top, width, height, data).chart
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(10)
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = y_title
    chart.value_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.tick_labels.font.size = Pt(9)
    chart.category_axis.tick_labels.font.size = Pt(9)
    for idx, s in enumerate(chart.series):
        s.format.line.color.rgb = PALETTE[idx]
        s.format.line.width = Pt(2.25)
        s.marker.format.fill.solid()
        s.marker.format.fill.fore_color.rgb = PALETTE[idx]
        s.marker.format.line.color.rgb = PALETTE[idx]
    return chart


def build_deck():
    regional = load_sheet('Regional_Performance')
    customer = load_sheet('Customer_Experience')
    risk = load_sheet('Operational_Risk')
    rp, cx, op = index_rows(regional), index_rows(customer), index_rows(risk)

    latest, previous, baseline = '2026-Q2', '2026-Q1', '2025-Q3'
    latest_revenue = sum(rp[(latest, r)]['Revenue_USD_M'] for r in REGION_ORDER)
    previous_revenue = sum(rp[(previous, r)]['Revenue_USD_M'] for r in REGION_ORDER)
    qoq_growth = pct_change(latest_revenue, previous_revenue)
    total_units = sum(rp[(latest, r)]['Units_Sold_K'] for r in REGION_ORDER)
    weighted_nps = sum(cx[(latest, r)]['NPS'] * rp[(latest, r)]['Units_Sold_K'] for r in REGION_ORDER) / total_units

    revenue_change = {r: pct_change(rp[(latest, r)]['Revenue_USD_M'], rp[(baseline, r)]['Revenue_USD_M']) for r in REGION_ORDER}
    best_growth_region = max(revenue_change, key=revenue_change.get)
    worst_growth_region = min(revenue_change, key=revenue_change.get)

    prs = Presentation(TEMPLATE_FILE)

    # Slide 1
    s = prs.slides[0]
    set_text(s, 'subtitle_1', 'Q2 2026 — Executive Review', 14, RGBColor(202, 220, 252), italic=True)
    set_text(s, 'kpi1_value', f'${latest_revenue:,.0f}M', 28, WHITE, bold=True, valign=MSO_ANCHOR.MIDDLE)
    set_text(s, 'kpi1_label', 'Q2 Total Revenue', 12, RGBColor(202, 220, 252), valign=MSO_ANCHOR.MIDDLE)
    set_text(s, 'kpi2_value', f'{qoq_growth:+.1f}%', 28, WHITE, bold=True, valign=MSO_ANCHOR.MIDDLE)
    set_text(s, 'kpi2_label', 'Quarter-over-Quarter Growth', 12, RGBColor(202, 220, 252), valign=MSO_ANCHOR.MIDDLE)
    set_text(s, 'kpi3_value', f'{weighted_nps:.1f}', 28, WHITE, bold=True, valign=MSO_ANCHOR.MIDDLE)
    set_text(s, 'kpi3_label', 'Unit-Weighted Customer NPS', 12, RGBColor(202, 220, 252), valign=MSO_ANCHOR.MIDDLE)
    summary = (
        f'Q2 revenue reached ${latest_revenue:,.0f}M, up {qoq_growth:.1f}% from Q1. '
        f'{best_growth_region} is the growth leader, with revenue up {revenue_change[best_growth_region]:.1f}% versus 2025-Q3, '
        f'while {worst_growth_region} declined {abs(revenue_change[worst_growth_region]):.1f}%. '
        f'The most urgent issue is Latin America, where NPS fell to {cx[(latest, "Latin America")]["NPS"]}, '
        f'stockouts rose to {op[(latest, "Latin America")]["Stockout_Rate_pct"]:.1f}%, and return processing reached '
        f'{op[(latest, "Latin America")]["Return_Processing_Days"]:.0f} days. APAC growth remains attractive, but '
        f'{op[(latest, "APAC")]["Warehouse_Capacity_Util_pct"]:.0f}% warehouse utilization creates a near-term capacity constraint.'
    )
    set_text(s, 'summary_text', summary, 16, WHITE, margin=0.02)
    replace_icon(s, 'icon_1', 'icon_1_label', 'growth_chart.png')

    # Slide 2
    s = prs.slides[1]
    set_text(s, 'title_2', 'APAC Accelerates While Latin America Contracts', 24, NAVY, bold=True, italic=False)
    rev_series = {r: [rp[(q, r)]['Revenue_USD_M'] for q in QUARTER_ORDER] for r in REGION_ORDER}
    add_line_chart(s, 'chart_performance', 'chart_performance_label', QUARTER_ORDER, rev_series, 'Revenue (USD M)')
    insights = [
        f'APAC revenue rose from ${rp[(baseline, "APAC")]["Revenue_USD_M"]:.0f}M to ${rp[(latest, "APAC")]["Revenue_USD_M"]:.0f}M, a {revenue_change["APAC"]:.1f}% increase.',
        f'North America remains the largest market at ${rp[(latest, "North America")]["Revenue_USD_M"]:.0f}M in Q2, with {rp[(latest, "North America")]["Units_Sold_K"]:,.0f}K units sold.',
        f'Latin America fell to ${rp[(latest, "Latin America")]["Revenue_USD_M"]:.0f}M and {rp[(latest, "Latin America")]["Units_Sold_K"]:,.0f}K units, both below 2025-Q3.',
        f'Latin America return rate increased to {rp[(latest, "Latin America")]["Returns_Rate_pct"]:.1f}%, more than double APAC’s {rp[(latest, "APAC")]["Returns_Rate_pct"]:.1f}%.'
    ]
    set_text(s, 'insights_performance', insights, 14, NAVY, bullet_lines=True, margin=0.02)
    replace_icon(s, 'icon_2', 'icon_2_label', 'growth_chart.png')

    # Slide 3
    s = prs.slides[2]
    set_text(s, 'title_3', 'Customer Experience Strengthens in APAC but Deteriorates in Latin America', 23, NAVY, bold=True, italic=False)
    nps_series = {r: [cx[(q, r)]['NPS'] for q in QUARTER_ORDER] for r in REGION_ORDER}
    add_line_chart(s, 'chart_customer', 'chart_customer_label', QUARTER_ORDER, nps_series, 'Net Promoter Score')
    insights = [
        f'APAC leads Q2 customer experience with NPS {cx[(latest, "APAC")]["NPS"]}, CSAT {cx[(latest, "APAC")]["CSAT_pct"]:.0f}%, and {cx[(latest, "APAC")]["Avg_Delivery_Days"]:.1f}-day delivery.',
        f'Europe improved NPS from {cx[(baseline, "Europe")]["NPS"]} to {cx[(latest, "Europe")]["NPS"]} while reducing delivery time to {cx[(latest, "Europe")]["Avg_Delivery_Days"]:.1f} days.',
        f'Latin America NPS dropped from {cx[(baseline, "Latin America")]["NPS"]} to {cx[(latest, "Latin America")]["NPS"]}, and CSAT declined to {cx[(latest, "Latin America")]["CSAT_pct"]:.0f}%.',
        f'Latin America support tickets climbed to {cx[(latest, "Latin America")]["Support_Tickets_K"]:.0f}K as delivery time reached {cx[(latest, "Latin America")]["Avg_Delivery_Days"]:.1f} days.'
    ]
    set_text(s, 'insights_customer', insights, 14, NAVY, bullet_lines=True, margin=0.02)
    replace_icon(s, 'icon_3', 'icon_3_label', 'customer_star.png')

    # Slide 4
    s = prs.slides[3]
    set_text(s, 'title_4', 'Latin America Risk Escalates as APAC Capacity Tightens', 24, NAVY, bold=True, italic=False)
    stockout_series = {r: [op[(q, r)]['Stockout_Rate_pct'] for q in QUARTER_ORDER] for r in REGION_ORDER}
    add_line_chart(s, 'chart_risk', 'chart_risk_label', QUARTER_ORDER, stockout_series, 'Stockout Rate (%)')
    insights = [
        f'Latin America stockouts reached {op[(latest, "Latin America")]["Stockout_Rate_pct"]:.1f}%, up from {op[(baseline, "Latin America")]["Stockout_Rate_pct"]:.1f}% in 2025-Q3.',
        f'Latin America return processing lengthened to {op[(latest, "Latin America")]["Return_Processing_Days"]:.0f} days and fraud incidents rose to {op[(latest, "Latin America")]["Fraud_Incidents"]:.0f}.',
        f'APAC warehouse utilization reached {op[(latest, "APAC")]["Warehouse_Capacity_Util_pct"]:.0f}% while stockouts increased to {op[(latest, "APAC")]["Stockout_Rate_pct"]:.1f}%.',
        f'North America and Europe remain comparatively stable, with Q2 stockout rates of {op[(latest, "North America")]["Stockout_Rate_pct"]:.1f}% and {op[(latest, "Europe")]["Stockout_Rate_pct"]:.1f}%.'
    ]
    set_text(s, 'insights_risk', insights, 14, NAVY, bullet_lines=True, margin=0.02)
    replace_icon(s, 'icon_4', 'icon_4_label', 'risk_alert.png')

    # Slide 5
    s = prs.slides[4]
    actions = [
        'STABILIZE Latin America fulfillment: set a weekly stockout-reduction plan focused on high-demand SKUs and supplier availability.',
        'REDESIGN Latin America returns and service operations: shorten the 10-day return cycle and address the drivers behind rising support tickets.',
        'EXPAND APAC capacity before further growth: protect service levels as warehouse utilization reaches 92% and stockouts continue to rise.',
        'SCALE the APAC playbook selectively: preserve customer experience while reallocating investment toward the strongest sustainable growth opportunities.'
    ]
    set_text(s, 'recommendations_text', actions, 18, WHITE, bullet_lines=True, margin=0.02)
    set_text(s, 'closing_statement', 'Leadership priority: resolve Latin America’s operating breakdown while creating capacity for APAC’s next phase of growth.', 15, RGBColor(202, 220, 252), italic=True, margin=0.02)
    replace_icon(s, 'icon_5', 'icon_5_label', 'lightbulb_action.png')

    prs.save(OUTPUT_FILE)
    print(f'Created: {OUTPUT_FILE}')


if __name__ == '__main__':
    build_deck()
